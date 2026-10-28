import io
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from typing import Optional
from app.database import get_db
from app import models, schemas
from app.workspace import get_workspace_user

router = APIRouter(prefix="/api/export", tags=["export"])

@router.get("/docx/{resume_id}")
def export_resume_docx(
    resume_id: int,
    db: Session = Depends(get_db)
):
    resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found.")

    doc = Document()

    # Document margins: 0.75 in
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    c = resume.content or {}
    p_info = c.get("personal_info", {})
    workspace_user = get_workspace_user(db)

    # Title / Name
    name_p = doc.add_paragraph()
    run_name = name_p.add_run(p_info.get("full_name") or workspace_user.full_name)
    run_name.font.size = Pt(20)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 23, 42)

    # Headline
    if p_info.get("headline"):
        h_p = doc.add_paragraph()
        run_h = h_p.add_run(p_info.get("headline"))
        run_h.font.size = Pt(12)
        run_h.font.color.rgb = RGBColor(79, 70, 229)

    # Contact line
    contact_parts = []
    if p_info.get("email"):
        contact_parts.append(p_info["email"])
    if p_info.get("phone"):
        contact_parts.append(p_info["phone"])
    if p_info.get("location"):
        contact_parts.append(p_info["location"])
    if p_info.get("linkedin"):
        contact_parts.append(p_info["linkedin"])
    if p_info.get("website"):
        contact_parts.append(p_info["website"])

    if contact_parts:
        c_p = doc.add_paragraph(" | ".join(contact_parts))
        c_p.paragraph_format.space_after = Pt(10)

    # Summary
    if c.get("summary"):
        s_heading = doc.add_paragraph()
        run_sh = s_heading.add_run("PROFESSIONAL SUMMARY")
        run_sh.font.bold = True
        run_sh.font.size = Pt(12)
        s_p = doc.add_paragraph(c["summary"])
        s_p.paragraph_format.space_after = Pt(10)

    # Experience
    if c.get("experience"):
        exp_h = doc.add_paragraph()
        run_eh = exp_h.add_run("EXPERIENCE")
        run_eh.font.bold = True
        run_eh.font.size = Pt(12)
        for exp in c["experience"]:
            role_p = doc.add_paragraph()
            role_run = role_p.add_run(f"{exp.get('role', '')} — {exp.get('company', '')}")
            role_run.font.bold = True
            date_line = f"{exp.get('start_date', '')} to {exp.get('end_date', '')}"
            if exp.get('location'):
                date_line += f" | {exp.get('location')}"
            d_p = doc.add_paragraph(date_line)
            d_p.paragraph_format.space_after = Pt(4)
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(bullet, style='List Bullet')
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Education
    if c.get("education"):
        edu_h = doc.add_paragraph()
        run_edh = edu_h.add_run("EDUCATION")
        run_edh.font.bold = True
        run_edh.font.size = Pt(12)
        for edu in c["education"]:
            deg_p = doc.add_paragraph()
            deg_run = deg_p.add_run(f"{edu.get('degree', '')} in {edu.get('field_of_study', '')}")
            deg_run.font.bold = True
            doc.add_paragraph(f"{edu.get('institution', '')} | Graduated: {edu.get('graduation_date', '')}")

    # Skills
    if c.get("skills"):
        sk_h = doc.add_paragraph()
        run_sk = sk_h.add_run("TECHNICAL SKILLS")
        run_sk.font.bold = True
        run_sk.font.size = Pt(12)
        for cat in c["skills"]:
            skills_str = ", ".join(cat.get("skills", []))
            sp = doc.add_paragraph()
            cat_run = sp.add_run(f"{cat.get('name', 'Skills')}: ")
            cat_run.font.bold = True
            sp.add_run(skills_str)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"{resume.title.replace(' ', '_')}.docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/json/{resume_id}")
def export_resume_json(
    resume_id: int,
    db: Session = Depends(get_db)
):
    resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found.")

    return {
        "version": "1.0",
        "exported_at": resume.updated_at.isoformat(),
        "title": resume.title,
        "template_id": resume.template_id,
        "theme_color": resume.theme_color,
        "font_family": resume.font_family,
        "font_size": resume.font_size,
        "spacing": resume.spacing,
        "target_role": resume.target_role,
        "ats_score": resume.ats_score,
        "content": resume.content
    }

@router.post("/import-json")
def import_resume_json(
    payload: dict,
    db: Session = Depends(get_db)
):
    workspace_user = get_workspace_user(db)
    content = payload.get("content") or {}
    title = payload.get("title") or "Imported Resume"
    template_id = payload.get("template_id") or "ats_classic"
    theme_color = payload.get("theme_color") or "#4F46E5"

    new_resume = models.Resume(
        user_id=workspace_user.id,
        title=f"{title} (Imported)",
        template_id=template_id,
        theme_color=theme_color,
        content=content
    )
    db.add(new_resume)
    db.commit()
    db.refresh(new_resume)

    return new_resume

@router.get("/account-data")
def export_all_account_data(
    db: Session = Depends(get_db)
):
    workspace_user = get_workspace_user(db)
    resumes = db.query(models.Resume).all()
    letters = db.query(models.CoverLetter).all()
    apps = db.query(models.JobApplication).all()

    return {
        "user": {
            "email": workspace_user.email,
            "full_name": workspace_user.full_name,
            "headline": workspace_user.headline,
            "theme_preference": workspace_user.theme_preference,
            "created_at": workspace_user.created_at.isoformat()
        },
        "resumes": [
            {
                "id": r.id,
                "title": r.title,
                "template_id": r.template_id,
                "theme_color": r.theme_color,
                "ats_score": r.ats_score,
                "content": r.content
            } for r in resumes
        ],
        "cover_letters": [
            {
                "id": l.id,
                "title": l.title,
                "company_name": l.company_name,
                "job_title": l.job_title,
                "content": l.content
            } for l in letters
        ],
        "job_applications": [
            {
                "id": a.id,
                "company": a.company,
                "job_title": a.job_title,
                "status": a.status,
                "notes": a.notes
            } for a in apps
        ]
    }
